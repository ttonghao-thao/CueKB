"""Canonical parsing and persistence invariants of the S1 source graph."""

import hashlib
import itertools
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError
from test_m3_postgres import pg as pg  # noqa: PLC0414 - pytest fixture re-export

from cuekb.api import v2
from cuekb.api.dependencies import current_principal
from cuekb.domain.document_graph import DocumentGraph
from cuekb.domain.models import KnowledgeBase
from cuekb.main import create_app
from cuekb.migrations import migrations, render_schema, upgrade
from cuekb.services.parsing import parse_document
from cuekb.worker import materialize_chunks


def test_tree_retains_raw_paragraph_and_distinct_repeated_headings():
    source = (
        "# Policy\n\n## Terms\n\n"
        + ("Long source sentence. " * 180)
        + "\n\n## Terms\n\nDifferent terms."
    )
    blocks = parse_document(source.encode(), "policy.md", "text/markdown", version_id=uuid4())
    graph = blocks.document_graph
    assert len(blocks) > 3
    sections = [n for n in graph.nodes if n.kind == "section" and n.text == "Terms"]
    assert len(sections) == 2 and sections[0].id != sections[1].id
    paragraphs = [n for n in graph.nodes if n.kind == "paragraph"]
    assert paragraphs[0].text == "Long source sentence. " * 179 + "Long source sentence."
    assert paragraphs[-1].parent_id == sections[-1].id
    anchor = paragraphs[0].source_spans[0].anchor
    assert source[anchor.start_offset : anchor.end_offset] == paragraphs[0].text
    changed = graph.model_dump()
    changed["nodes"][1]["parent_id"] = changed["nodes"][-1]["id"]
    with pytest.raises(ValidationError):
        DocumentGraph.model_validate(changed)


def test_table_coordinates_and_explicit_footnote_are_not_neighborhood_guesses():
    source = "# Specifications\n\n| Model | Voltage |\n| --- | --- |\n| A100 | 12V |\n| A200 | 24V |\n\nOnly sealed packages.[^storage]\n\n[^storage]: Store below 25 C.\n"
    graph = parse_document(source.encode(), "table.md", "text/markdown").document_graph
    assert any(e.type == "header_for" and e.required for e in graph.edges)
    assert len([n for n in graph.nodes if n.kind == "cell"]) == 6
    assert any(e.type == "footnote_for" and e.required for e in graph.edges)
    assert all(
        e.required is False for e in graph.edges if e.type in {"parent_of", "next_in_structure"}
    )
    missing = parse_document(b"Condition[^missing].", "x.txt", None).document_graph
    assert missing.status == "limited" and "unresolved_footnote" in missing.issues


def test_general_document_publish_graph_and_acl(pg, monkeypatch):
    _, _, key = pg.create_api_key("graph-owner", "graph-test")
    principal = pg.authenticate(key)
    kb = pg.create_knowledge_base(KnowledgeBase(name="policies"), principal)
    source = b"# Leave policy\n\nEmployees must apply before travel.\n\n## Conditions\n\nApproval is required."
    job = pg.create_ingestion_job(
        principal_id=principal.id,
        kb_id=kb.id,
        name="Policy",
        source_uri="test://original",
        filename="policy.md",
        media_type="text/markdown",
        content_sha256=hashlib.sha256(source).hexdigest(),
        business_version=None,
        scope={},
        idempotency_key="graph-ingest",
        request_sha256="f" * 64,
        document_id=None,
        auto_publish=True,
        processing_snapshot={
            "content_kind": "general",
            "source_metadata": {"source_owner": "HR", "language": "en"},
        },
    )
    claimed = pg.claim_job("graph-test", 300, 3)
    chunks = materialize_chunks(
        pg, SimpleNamespace(read=lambda _, source=source: source), claimed, template="evidence-v2"
    )
    assert chunks.document_graph and len(chunks.document_graph.units) == len(chunks)
    assert all(c.evidence_structure.fact_kind == "general" for c in chunks)
    assert pg.save_ready(job.id, "graph-test", chunks)
    with pg.read_snapshot() as snapshot:
        authority = snapshot.evidence_authority_snapshot(
            principal, [kb.id], [chunks[0].id]
        )
    assert authority["chunks"][0].source_text == chunks[0].source_text
    assert authority["node_ids"][chunks[0].id]
    assert authority["content_revisions"][str(kb.id)] >= 1
    assert authority["authorization_revisions"][str(kb.id)] >= 0
    with pg._read() as conn:
        saved = pg._load_document_graph(conn, job.version_id)
        assert saved.digest() == chunks.document_graph.digest()
        assert (
            conn.execute(
                text("SELECT graph_digest FROM publication_snapshots WHERE version_id=:v"),
                {"v": job.version_id},
            ).scalar()
            == saved.digest()
        )
    page = pg.document_graph_page(job.document_id, job.version_id, principal, limit=2)
    assert (
        page["total"] > 2 and len(page["nodes"]) == 2 and page["metadata"]["source_owner"] == "HR"
    )
    monkeypatch.setattr(v2, "repository", lambda: pg)
    app = create_app()
    app.dependency_overrides[current_principal] = lambda: principal
    path = f"/v2/documents/{job.document_id}/versions/{job.version_id}/structure"
    with TestClient(app) as client:
        assert client.get(path).status_code == 200
        assert client.get(path + "?limit=101").status_code == 422
        pg.delete_document(job.document_id)
        assert client.get(path).status_code == 404


def test_legacy_schema_upgrade_preserves_documents_and_checksums(pg):
    kb = pg.create_knowledge_base(KnowledgeBase(name="preserve"))
    with pg._write() as conn:
        for table in [
            "scope_assertions",
            "publication_snapshots",
            "retrieval_unit_nodes",
            "retrieval_units",
            "evidence_edges",
            "document_nodes",
            "document_graphs",
            "schema_revisions",
        ]:
            conn.execute(text(f"DROP TABLE {table} CASCADE"))
    result = upgrade(pg.engine)
    assert result["revision"] == "0003" and pg.get_knowledge_base(kb.id).name == "preserve"
    assert upgrade(pg.engine) == result
    with pg._write() as conn:
        conn.execute(
            text("UPDATE schema_revisions SET checksum=repeat('0',64) WHERE revision='0001'")
        )
    with pytest.raises(RuntimeError, match="checksum"):
        upgrade(pg.engine)


def test_generated_empty_schema_is_the_migration_chain():
    assert Path("db/schema.sql").read_text() == render_schema()
    assert [r for r, _ in migrations()] == ["0001", "0002", "0003"]


def test_review_derivation_preserves_source_tree_without_inventing_removed_facts(pg):
    from test_choice_review import finish, setup

    from cuekb.schemas import ReviewDecisionsRequest

    source = b"# Policy\n\nFirst instruction.\n\nSecond instruction."
    principal, _, _, job = setup(pg, source)
    chunks = finish(pg, job, source)
    plan = pg.create_review_plan(job.document_id, job.version_id, principal)
    plan = pg.decide_review(
        plan["id"],
        principal,
        ReviewDecisionsRequest(
            revision=plan["revision"],
            decisions=[
                {"fact_id": chunks[0].id, "action": "general"},
                {"fact_id": chunks[1].id, "action": "exclude"},
            ],
        ),
    )
    derived = pg.apply_review_plan(plan["id"], principal, plan["revision"])
    new = finish(pg, derived, source)
    assert len(new) == 1 and derived.version_id == job.version_id
    assert new.document_graph.source_version_id is None
    assert any(n.text == "Second instruction." for n in new.document_graph.nodes)
    assert len(new.document_graph.units) == 1
    assert new.document_graph.units[0].scope["source_mapping"] == "resolved"
    with pg._read() as conn:
        old = pg._load_document_graph(conn, job.version_id)
        assert {n.id for n in old.nodes} == {n.id for n in chunks.document_graph.nodes}
    # Restoring a version retains its canonical structure in a new identity space.
    pg.publish(job.document_id, derived.version_id, "default", principal)
    restored = pg.create_restore_job(
        principal_id=principal.id,
        document_id=job.document_id,
        source_version_id=derived.version_id,
        idempotency_key="restore-graph",
        request_sha256="d" * 64,
    )
    restored_chunks = finish(pg, restored, source)
    assert restored_chunks.document_graph.source_version_id == derived.version_id
    key_id, reader_id, key = pg.create_api_key("graph-reader", "local")
    reader = pg.authenticate(key)
    pg.grant(job.kb_id, reader_id, "read")
    with pytest.raises(PermissionError):
        pg.document_graph_page(job.document_id, restored.version_id, reader)
    assert (
        pg.document_graph_page(job.document_id, derived.version_id, reader)["status"] == "complete"
    )
    pg.revoke_api_key(key_id)
    with pytest.raises(PermissionError):
        pg.document_graph_page(job.document_id, derived.version_id, reader)


def test_footnote_dependency_is_delivered_and_missing_definition_blocks_publication(pg):
    _, _, key = pg.create_api_key("footnotes", "local")
    principal = pg.authenticate(key)
    kb = pg.create_knowledge_base(KnowledgeBase(name="footnotes"), principal)
    for index, source in enumerate(
        [
            b"# Rules\n\nValid only as described[^a].\n\n[^a]: Keep sealed.",
            b"# Rules\n\nValid only as described[^missing].",
        ]
    ):
        job = pg.create_ingestion_job(
            principal_id=principal.id,
            kb_id=kb.id,
            name="Footnotes",
            source_uri="test://fn",
            filename="footnotes.md",
            media_type="text/markdown",
            content_sha256=hashlib.sha256(source).hexdigest(),
            business_version=None,
            scope={},
            idempotency_key=f"fn-{index}",
            request_sha256=str(index) * 64,
            document_id=None,
            auto_publish=True,
            processing_snapshot={"content_kind": "general"},
        )
        claimed = pg.claim_job("fn", 300, 3)
        chunks = materialize_chunks(
            pg,
            SimpleNamespace(read=lambda _, source=source: source),
            claimed,
            template="evidence-v2",
        )
        published = pg.save_ready(job.id, "fn", chunks)
        if index == 0:
            assert published and chunks[0].evidence_structure.context_links
            assert chunks.document_graph.dependencies
        else:
            assert not published and pg.get_job(job.id).stage == "needs_review"


def test_source_metadata_validation():
    from cuekb.schemas import V2DocumentMetadata

    with pytest.raises(ValidationError):
        V2DocumentMetadata(
            kb_id=uuid4(), name="bad dates", source_metadata={"valid_from": "2026-10-01T00:00:00"}
        )
    with pytest.raises(ValidationError):
        V2DocumentMetadata(
            kb_id=uuid4(),
            name="bad range",
            source_metadata={
                "valid_from": "2026-10-02T00:00:00Z",
                "valid_to": "2026-10-01T00:00:00Z",
            },
        )


def test_real_docx_keeps_table_structure_on_cpu(tmp_path):
    import io

    pytest.importorskip("docling")
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_heading("Policy", level=1)
    document.add_paragraph("Requests need prior approval.")
    document.add_heading("Approval limits", level=2)
    document.add_paragraph("Only managers can approve exceptions.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Approval"
    table.cell(1, 0).text = "Staff"
    table.cell(1, 1).text = "Manager"
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = parse_document(buffer.getvalue(), "policy.docx", None, version_id=uuid4())
    nested = next(
        n
        for n in parsed.document_graph.nodes
        if n.kind == "section" and n.text == "Approval limits"
    )
    parent = next(n for n in parsed.document_graph.nodes if n.id == nested.parent_id)
    assert parent.kind == "section" and parent.text == "Policy"
    assert any(n.kind == "table" for n in parsed.document_graph.nodes)
    assert any(e.type == "header_for" for e in parsed.document_graph.edges)
    assert any("Requests need prior approval." in n.text for n in parsed.document_graph.nodes)


def test_migrations_empty_database_and_atomic_failure(pg, monkeypatch):
    import cuekb.migrations as manager

    with pg._write() as conn:
        schema = conn.execute(text("SELECT current_schema()")).scalar()
        assert schema.startswith("cuekb_test_") and schema.replace("_", "").isalnum()
        conn.exec_driver_sql(f"DROP SCHEMA {schema} CASCADE")
        conn.exec_driver_sql(f"CREATE SCHEMA {schema}")
    original = manager.migrations
    monkeypatch.setattr(
        manager,
        "migrations",
        lambda: [
            *original(),
            (
                "0003",
                "CREATE TABLE partial_write(id integer); SELECT intentional_migration_failure();",
            ),
        ],
    )
    with pytest.raises(DBAPIError):
        manager.upgrade(pg.engine)
    with pg._read() as conn:
        assert conn.execute(text("SELECT to_regclass('schema_revisions')")).scalar() is None
        assert conn.execute(text("SELECT to_regclass('partial_write')")).scalar() is None
    monkeypatch.setattr(manager, "migrations", original)
    assert manager.upgrade(pg.engine)["revision"] == "0003"
    with pg._read() as conn:
        assert manager.check_connection(conn)["revision"] == "0003"


def test_required_closure_is_bounded_and_reported():
    from cuekb.domain.document_graph import GraphBuilder, compute_dependencies

    b = GraphBuilder(uuid4(), "a" * 64)
    nodes = [b.add("paragraph", b.root, str(i), [], [], str(i)) for i in range(5)]
    for a, c in itertools.pairwise(nodes):
        b.edge(a, c, "governed_by", {"kind": "explicit"}, True)
    graph = b.finish()
    compute_dependencies(graph)
    assert graph.status == "limited"
    assert graph.dependencies[str(nodes[0])]["state"] == "limited"


def test_s1_validation_fixtures_are_real_files_and_report_uses_actual_worker_state():
    import io

    import httpx

    from cuekb.validation_s1 import fixtures, validate

    pdfium = pytest.importorskip("pypdfium2")
    source = fixtures()
    pdf = pdfium.PdfDocument(source["pdf"])
    try:
        assert "Approval is required before travel." in pdf[0].get_textpage().get_text_range()
    finally:
        pdf.close()
    from docx import Document

    doc = Document(io.BytesIO(source["docx"]))
    assert doc.tables[0].cell(1, 1).text == "Manager"

    def unavailable(request):
        if request.url.path == "/v2/ready":
            return httpx.Response(200, json={"status": "ready"})
        return httpx.Response(
            503, json={"error_code": "MODEL_UNAVAILABLE", "trace_id": "bounded-trace"}
        )

    with httpx.Client(
        base_url="http://fixture", transport=httpx.MockTransport(unavailable)
    ) as client:
        report = validate(client, uuid4(), ["txt"], 1)
    assert (
        report["passed"] == 0
        and report["failed"] == 1
        and "MODEL_UNAVAILABLE" in report["cases"][0]["error"]
    )


def test_upgrade_requires_drained_jobs_and_does_not_bless_partial_schema(pg):
    from test_choice_review import setup

    _, _, _, job = setup(pg, b"Pending instruction.")
    with pg._write() as conn:
        for table in [
            "scope_assertions",
            "publication_snapshots",
            "retrieval_unit_nodes",
            "retrieval_units",
            "evidence_edges",
            "document_nodes",
            "document_graphs",
            "schema_revisions",
        ]:
            conn.execute(text(f"DROP TABLE {table} CASCADE"))
    with pytest.raises(RuntimeError, match="drained"):
        upgrade(pg.engine)
    with pg._read() as conn:
        assert conn.execute(text("SELECT to_regclass('schema_revisions')")).scalar() is None
    with pg._write() as conn:
        conn.execute(text("UPDATE ingestion_jobs SET status='failed' WHERE id=:id"), {"id": job.id})
        conn.execute(text("DROP INDEX document_versions_content_processing_unique"))
    with pytest.raises(RuntimeError, match="unsupported_legacy_schema"):
        upgrade(pg.engine)


def test_readiness_reports_schema_upgrade_required(pg, monkeypatch):
    from cuekb.api import routes

    with pg._write() as conn:
        conn.execute(text("DROP TABLE schema_revisions"))
    monkeypatch.setattr(routes, "repository", lambda: pg)
    monkeypatch.setattr(routes, "get_settings", lambda: SimpleNamespace(backend="production"))
    monkeypatch.setattr(
        routes,
        "search_backend",
        lambda: SimpleNamespace(
            client=SimpleNamespace(cluster=SimpleNamespace(health=lambda: {"status": "green"}))
        ),
    )
    with TestClient(create_app()) as client:
        response = client.get("/v2/ready")
    assert (
        response.status_code == 503 and response.json()["error_code"] == "SCHEMA_UPGRADE_REQUIRED"
    )


def test_validation_report_runs_real_upload_parse_and_publish(pg, monkeypatch, tmp_path):
    import httpx

    import cuekb.validation_s1 as validation
    from cuekb.adapters.storage import LocalFileStorage
    from cuekb.api import routes

    _, _, key = pg.create_api_key("validation-owner", "test")
    principal = pg.authenticate(key)
    kb = pg.create_knowledge_base(KnowledgeBase(name="validation"), principal)
    pg.save_model_configuration("http://local-fixture", "local-key", "test-embedding", "", "", "")
    storage = LocalFileStorage(str(tmp_path))
    monkeypatch.setattr(routes, "repository", lambda: pg)
    monkeypatch.setattr(v2, "repository", lambda: pg)
    monkeypatch.setattr(routes, "file_storage", lambda: storage)
    monkeypatch.setattr(validation.time, "sleep", lambda _: None)
    app = create_app()
    app.dependency_overrides[current_principal] = lambda: principal
    with TestClient(app) as api:

        def bridge(request):
            response = api.request(
                request.method,
                str(request.url.copy_with(scheme="http", host="testserver")),
                headers=dict(request.headers),
                content=request.read(),
            )
            if (
                request.method == "POST"
                and request.url.path == "/v2/documents"
                and response.status_code == 202
            ):
                job = pg.claim_job("validator", 300, 3)
                chunks = materialize_chunks(pg, storage, job, template="evidence-v2")
                pg.save_ready(job["id"], "validator", chunks)
            return httpx.Response(
                response.status_code, content=response.content, headers=dict(response.headers)
            )

        with httpx.Client(
            base_url="http://testserver", transport=httpx.MockTransport(bridge)
        ) as client:
            report = validation.validate(client, kb.id, ["txt", "md", "docx"], 10)
    assert report["passed"] == 3 and report["failed"] == 0
    assert all(case["published_version_id"] == case["version_id"] for case in report["cases"])


def test_worker_refuses_old_schema_before_claiming_jobs(pg, monkeypatch):
    from cuekb import worker

    with pg._write() as conn:
        conn.execute(text("DROP TABLE schema_revisions"))
    monkeypatch.setattr(worker, "get_settings", lambda: SimpleNamespace(backend="production"))
    monkeypatch.setattr(worker, "runtime", lambda: SimpleNamespace(repository=pg))
    with pytest.raises(RuntimeError, match="schema_upgrade_required"):
        worker.main()


def test_long_source_is_preserved_while_structure_previews_are_bounded(pg):
    from test_choice_review import finish, setup

    source = ("# Policy\n\n" + "Long source word. " * 2400).encode()
    principal, _, _, job = setup(pg, source)
    chunks = finish(pg, job, source)
    paragraph = next(n for n in chunks.document_graph.nodes if n.kind == "paragraph")
    assert len(paragraph.text) > 32768 and paragraph.parse_quality == "limited"
    span = paragraph.source_spans[0]
    raw = source.decode()
    assert raw[span.anchor.start_offset : span.anchor.end_offset] == span.text
    page = pg.document_graph_page(job.document_id, job.version_id, principal)
    node = next(n for n in page["nodes"] if n["kind"] == "paragraph")
    assert len(node["text"]) == 10000 and node["text_truncated"]
    assert node["source_spans"][0]["text_truncated"] and len(node["source_spans"][0]["text"]) == 500
